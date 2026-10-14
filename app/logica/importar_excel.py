"""Importar Excel (v2, Módulo K): leer la planilla y detectar sus filas.

Reglas: docs/requisitos.md §4.12 (RN-60 a RN-65). Este archivo solo LEE el Excel:
no guarda nada en la base. La lista de revisión y el guardado vienen después.

Se trabaja con el contenido del archivo (bytes) y no con su ruta: así sirve
igual en la ventana de escritorio y en el modo Edge, donde el navegador entrega
el archivo y no su ruta. El archivo nunca se modifica ni se copia (RN-60).
"""

import io
import math
import re
from dataclasses import dataclass, field, replace
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.logica.errores import ErrorDeValidacion
from app.logica.formato import clave_orden

MENSAJE_ARCHIVO_INVALIDO = (
    "No se pudo leer el archivo. Solo se aceptan archivos .xlsx de Excel."
)

# Encabezados que marcan una columna de pagado o de cobrado (RN-64), ya normalizados.
TITULOS_PAGADO = {"pagado", "pago", "pagos", "abonado"}
TITULO_COBRADO = "cobrado"

# "82.000", "1.050.000,5" (punto cada 3 dígitos) o "260000", "41116,66" (sin puntos).
_NUMERO_ARGENTINO = re.compile(r"(\d{1,3}(?:\.\d{3})+|\d+)(?:,(\d+))?")


def normalizar(texto: str) -> str:
    """RN-67: sin mayúsculas, sin acentos y sin espacios de más."""
    return " ".join(clave_orden(texto).split())


# ---------------------------------------------------------------------------
# Celdas (RN-61, RN-62)
# ---------------------------------------------------------------------------


def _a_centavos(valor: Decimal) -> int:
    """Redondea al centavo, mitad hacia arriba (RN-03): 22.093,666 -> 2209367."""
    return int((valor * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def leer_numero_texto(texto: str) -> int | None:
    """RN-62: un texto en formato argentino, en centavos; None si no es un número.

    Acepta un "$" y espacios adelante, y un "-" al principio: "$ 82.000",
    "-5.000", "66.020,12". El punto vale solo como separador de miles de a
    3 dígitos: "82.5" no es un número (es texto).
    """
    resto = texto.strip()
    negativo = resto.startswith("-")
    if negativo:
        resto = resto[1:].lstrip()
    if resto.startswith("$"):
        resto = resto[1:].lstrip()
    # fullmatch: el texto entero tiene que ser el número ("3/3" o "US$ 20" no).
    encontrado = _NUMERO_ARGENTINO.fullmatch(resto)
    if not encontrado:
        return None
    enteros = encontrado.group(1).replace(".", "")
    decimales = encontrado.group(2) or "0"
    centavos = _a_centavos(Decimal(f"{enteros}.{decimales}"))
    return -centavos if negativo else centavos


@dataclass(frozen=True)
class Celda:
    """Lo que vale una celda según RN-61: un texto, un número (en centavos) o nada."""

    texto: str | None = None
    centavos: int | None = None
    es_suma: bool = False  # Viene de una fórmula SUM (RN-66).

    @property
    def vacia(self) -> bool:
        return self.texto is None and self.centavos is None


VACIA = Celda()


def clasificar(valor, tipo_de_dato: str = "") -> Celda:
    """RN-61: qué cuenta como número, texto o vacío.

    tipo_de_dato es el de openpyxl ("e" = error de Excel, como #DIV/0!).
    """
    if valor is None or tipo_de_dato == "e":
        return VACIA
    # bool va antes que int: en Python True y False también son números (1 y 0).
    if isinstance(valor, (bool, datetime, date, time, timedelta)):
        return VACIA
    if isinstance(valor, float) and not math.isfinite(valor):
        return VACIA  # Infinito o "no es un número": no es un monto (BUG-28).
    if isinstance(valor, (int, float)):
        # str() evita el error de los decimales binarios: 22093.666 queda exacto.
        return Celda(centavos=_a_centavos(Decimal(str(valor))))
    texto = str(valor).strip()
    if not texto:
        return VACIA
    centavos = leer_numero_texto(texto)
    if centavos is not None:
        return Celda(centavos=centavos)
    return Celda(texto=texto)


# ---------------------------------------------------------------------------
# Filas detectadas (RN-63, RN-64, RN-65)
# ---------------------------------------------------------------------------


@dataclass
class FilaDetectada:
    """Una fila de la lista de revisión tal como sale del Excel.

    fila_excel/columna: dónde está el nombre (o el número, si no tiene nombre).
    cobrado: el valor crudo de la columna "Cobrado" (centavos o texto); cómo se
    interpreta depende del tipo de la fila y se resuelve en la revisión.
    """

    fila_excel: int
    columna: int
    nombre: str | None = None
    monto_cent: int | None = None
    pagado_cent: int | None = None
    cobrado: int | str | None = None
    bloque: str | None = None
    es_titulo: bool = False
    # El monto sale de una fórmula de suma (SUM): es un total (RN-66).
    monto_es_suma: bool = False

    @property
    def completa(self) -> bool:
        """Con nombre y monto va a la lista principal; si no, a "Filas incompletas"."""
        return self.nombre is not None and self.monto_cent is not None


@dataclass
class HojaLeida:
    nombre: str
    filas: list[FilaDetectada] = field(default_factory=list)
    formulas_sin_valor: int = 0
    ejemplo_formula: str | None = None  # Ej. "D9", para el aviso de RN-61.

    @property
    def sin_datos(self) -> bool:
        return not self.filas

    @property
    def sin_filas_completas(self) -> bool:
        return not any(f.completa for f in self.filas)


def _tipo_de_encabezado(texto: str) -> str | None:
    clave = normalizar(texto)
    if clave in TITULOS_PAGADO:
        return "pagado"
    if clave == TITULO_COBRADO:
        return "cobrado"
    return None


def _detectar_filas(grilla: list[list[Celda]]) -> list[FilaDetectada]:
    """RN-63 y RN-64: recorre la grilla de arriba abajo y de izquierda a derecha."""
    filas: list[FilaDetectada] = []
    # Para cada columna, de qué tipo es según su último encabezado: "pagado",
    # "cobrado" o None. El encabezado es el último texto de arriba que no fue
    # el valor de una fila (RN-64).
    tipo_columna: dict[int, str | None] = {}
    for numero_fila, celdas in enumerate(grilla, start=1):
        if all(c.vacia for c in celdas):
            # Una fila totalmente vacía cierra todos los encabezados, como cierra
            # los bloques (RN-65): un "Pagado" de arriba no se estira a tablas de
            # más abajo que tienen el monto en esa columna (O63).
            tipo_columna.clear()
            continue
        abierta: FilaDetectada | None = None  # La fila que recibe los números.
        for numero_col, celda in enumerate(celdas, start=1):
            tipo = tipo_columna.get(numero_col)
            if celda.texto is not None:
                if tipo and abierta and abierta.completa:
                    # Texto al lado de una fila con nombre y monto: es su valor.
                    # En pagado se descarta (tiene que ser un número); en
                    # cobrado se guarda para leerlo en la revisión.
                    if tipo == "cobrado" and abierta.cobrado is None:
                        abierta.cobrado = celda.texto
                    continue
                # Cualquier otro texto es encabezado de las celdas de abajo y
                # empieza una fila nueva (RN-63 paso 1).
                tipo_columna[numero_col] = _tipo_de_encabezado(celda.texto)
                abierta = FilaDetectada(numero_fila, numero_col, nombre=celda.texto)
                filas.append(abierta)
            elif celda.centavos is not None:
                if tipo == "pagado":
                    # Nunca empieza una fila: es el pagado de la fila abierta.
                    if abierta and abierta.nombre and abierta.pagado_cent is None:
                        abierta.pagado_cent = celda.centavos
                elif tipo == "cobrado":
                    if abierta and abierta.nombre and abierta.cobrado is None:
                        abierta.cobrado = celda.centavos
                elif abierta and abierta.nombre and abierta.monto_cent is None:
                    abierta.monto_cent = celda.centavos  # RN-63 paso 2.
                    abierta.monto_es_suma = celda.es_suma
                elif abierta and abierta.nombre:
                    pass  # Un segundo número (ej. "Total deudas"): se descarta.
                else:
                    # RN-63 paso 3: número sin texto antes -> fila sin nombre.
                    # No "abre" la fila: los números siguientes no se le suman.
                    filas.append(
                        FilaDetectada(numero_fila, numero_col, monto_cent=celda.centavos)
                    )
    for fila in filas:
        if fila.monto_cent is None:
            # Un pagado sin monto no se usa: si después el usuario completa el
            # monto, la fila queda sin pagado (O62, RN-64).
            fila.pagado_cent = None
    return filas


def _asignar_bloques(
    filas: list[FilaDetectada],
    grilla: list[list[Celda]],
    combinadas: dict[tuple[int, int], int],
) -> None:
    """RN-65: marca los títulos y les asigna a las filas el título de su bloque.

    combinadas: para un título en una celda combinada, (fila, columna) -> la
    última columna que cubre.
    """
    nombres = {(f.fila_excel, f.columna): f for f in filas if f.nombre is not None}

    for fila in nombres.values():
        # Pegada: justo arriba, en la misma columna, hay el nombre de otra fila.
        pegada = (fila.fila_excel - 1, fila.columna) in nombres
        fila.es_titulo = fila.monto_cent is None and not pegada

    def fila_vacia(numero_fila: int) -> bool:
        return all(c.vacia for c in grilla[numero_fila - 1])

    for titulo in [f for f in nombres.values() if f.es_titulo]:
        ultima_col = combinadas.get((titulo.fila_excel, titulo.columna), titulo.columna)
        columnas = range(titulo.columna, ultima_col + 1)
        for numero_fila in range(titulo.fila_excel + 1, len(grilla) + 1):
            if fila_vacia(numero_fila):
                break
            debajo = [
                nombres[(numero_fila, c)] for c in columnas if (numero_fila, c) in nombres
            ]
            if any(f.es_titulo for f in debajo):
                break  # El título siguiente en esa columna.
            for f in debajo:
                if f.bloque is None:
                    f.bloque = titulo.nombre


# ---------------------------------------------------------------------------
# Archivo y hojas (RN-60, RN-61)
# ---------------------------------------------------------------------------


class LibroExcel:
    """Un .xlsx abierto para importar. Se crea con abrir_libro()."""

    def __init__(self, contenido: bytes) -> None:
        # Se abre dos veces: con data_only=True se ven los valores que Excel
        # guardó; con data_only=False se ven las fórmulas, para saber cuáles no
        # tienen valor guardado (RN-61).
        self._valores = load_workbook(io.BytesIO(contenido), data_only=True)
        self._formulas = load_workbook(io.BytesIO(contenido), data_only=False)

    def hojas(self) -> list[str]:
        """Las hojas que tienen alguna celda con algo (RN-60). Las hojas que son
        solo un gráfico no están en worksheets, así que no aparecen.

        Se mira el libro de fórmulas: una hoja con solo fórmulas sin calcular
        tiene celdas, aunque sus valores estén vacíos (BUG-27)."""
        return [
            hoja.title
            for hoja in self._formulas.worksheets
            if any(
                celda.value is not None
                for fila in hoja.iter_rows()
                for celda in fila
            )
        ]

    def leer_hoja(self, nombre: str) -> HojaLeida:
        hoja = self._valores[nombre]
        hoja_formulas = self._formulas[nombre]
        resultado = HojaLeida(nombre)

        grilla: list[list[Celda]] = []
        # Las dos hojas tienen las mismas filas y columnas: se recorren a la par.
        # (Pedir hoja_formulas.cell() por cada celda vacía crea celdas nuevas y es
        # lento con hojas grandes, O58.)
        # Se recorre solo hasta la última fila y columna con algo escrito. Las
        # celdas vacías con formato (un relleno en la fila 100.000) agrandan la
        # hoja y harían crear millones de celdas vacías (O58). _cells es el
        # diccionario interno de openpyxl con las celdas que existen en el archivo.
        con_valor = [
            coordenada
            for coordenada, celda in hoja_formulas._cells.items()
            if celda.value is not None
        ]
        ultima_fila = max((f for f, _ in con_valor), default=0)
        ultima_col = max((c for _, c in con_valor), default=0)
        if ultima_fila == 0:
            # Sin ningún valor. Ojo: iter_rows(max_row=0) no recorre "cero filas",
            # sino toda la hoja (para openpyxl, 0 es "sin límite", BUG-29).
            filas_valores = filas_formulas = []
        else:
            filas_valores = hoja.iter_rows(
                min_row=1, min_col=1, max_row=ultima_fila, max_col=ultima_col
            )
            filas_formulas = hoja_formulas.iter_rows(
                min_row=1, min_col=1, max_row=ultima_fila, max_col=ultima_col
            )
        for fila, fila_formulas in zip(filas_valores, filas_formulas, strict=True):
            renglon = []
            for celda, con_formula in zip(fila, fila_formulas, strict=True):
                valor = celda.value
                # data_type "n": sin valor de ningún tipo. Una fórmula que Excel
                # calculó y dio texto vacío ("") también llega como None, pero
                # con data_type "str": esa sí tiene valor guardado (BUG-26).
                if valor is None and celda.data_type == "n":
                    # ¿Fórmula sin valor guardado? (archivo nunca guardado en Excel)
                    if con_formula.data_type == "f":
                        resultado.formulas_sin_valor += 1
                        if resultado.ejemplo_formula is None:
                            resultado.ejemplo_formula = (
                                f"{get_column_letter(celda.column)}{celda.row}"
                            )
                # Las celdas combinadas que no son la primera vienen con valor
                # None: cuentan como vacías (RN-61).
                clasificada = clasificar(valor, celda.data_type)
                # Excel guarda las funciones en inglés: =SUMA(...) queda "SUM(".
                if con_formula.data_type == "f" and "SUM(" in str(con_formula.value).upper():
                    clasificada = replace(clasificada, es_suma=True)
                renglon.append(clasificada)
            grilla.append(renglon)

        combinadas = {
            (rango.min_row, rango.min_col): rango.max_col
            for rango in hoja.merged_cells.ranges
        }
        resultado.filas = _detectar_filas(grilla)
        _asignar_bloques(resultado.filas, grilla, combinadas)
        return resultado


def abrir_libro(contenido: bytes) -> LibroExcel:
    """RN-60: abre el .xlsx. Si no se puede leer, ErrorDeValidacion con el mensaje
    para el usuario. Quien llama ya controló que el nombre termine en .xlsx."""
    try:
        return LibroExcel(contenido)
    except Exception as error:
        # Se atrapa cualquier error a propósito (BUG-25): un archivo que no es
        # un .xlsx (un .txt renombrado, un .xls viejo, uno con contraseña) o un
        # .xlsx dañado por dentro hace fallar a openpyxl de muchas formas
        # distintas (zip, XML, compresión, estilos…). Para el usuario todas son
        # lo mismo: el archivo no se puede leer. "from error" guarda el detalle.
        raise ErrorDeValidacion(MENSAJE_ARCHIVO_INVALIDO) from error


def es_nombre_xlsx(nombre_archivo: str) -> bool:
    """RN-60: solo .xlsx (no .xls, .xlsm, .csv, .ods…)."""
    return nombre_archivo.lower().endswith(".xlsx")
