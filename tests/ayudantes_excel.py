"""Ayudantes para armar archivos .xlsx inventados en las pruebas de importar_excel.

Nunca se usa un archivo de datos-privados/: todo se genera acá con openpyxl.
Los archivos se arman en memoria (bytes), igual que los recibe la lógica.
"""

import io
import re
import zipfile

from openpyxl import Workbook


def celdas_planilla():
    """Las celdas de `planilla-prueba.xlsx` (tabla al principio de los CA del Módulo K).

    Devuelve un diccionario {"C3": "Alquiler", ...}. Los números de Excel son int/float;
    los "números escritos como texto" son str ("82.000", "66.020,12").
    """
    return {
        "C1": "Gasto del Mes",
        "D2": "Gastos",
        "E2": "Pagado",
        "C3": "Alquiler",
        "D3": 260000,
        "C4": "Servicios",
        "D4": "82.000",  # texto
        "C5": "Internet",
        "D5": 38700,
        "E5": 38700,
        "C6": "Tarjetas de Credito",
        "D6": 582917,
        "C7": "psicologo",
        "D7": 140000,
        "E7": 60000,
        "D8": "total",
        "E8": "Falta por pagar",
        "D9": 1103617,
        "E9": 1004917,
        "C11": "Deudas",
        "D11": "Monto",
        "E11": "Total deudas",
        "C12": "Monotributo 66.020,12",
        "D12": "66.020,12",
        "E12": 66021,  # D12 texto
        "C13": "Avant Service cuota 3/ 3 $41.116,66",
        "D13": 41116.66,
        "C14": "Deuda freezer",
        "C15": "Gym",
        "D15": 50000,
        "C16": "Obra social",
        "B18": "Sueldo",
        "B19": 1300,
        "D21": "Deudas sofia",
        "D22": "Prestamo MP",
        "E22": 20000,
        "D23": "Internet",
        "E23": 12900,
        "D24": "Total",
        "E24": 32900,
        "D26": "Marsu",
        "D27": "tacones chiqui baby",
        "E27": 6785,
        "D28": "Total",
        "E28": 6785,
    }


def armar_libro(hojas, combinadas=None, filas_ocultas=None):
    """Arma un Workbook de openpyxl.

    hojas: {"Octubre": {"C3": "Luz", ...}, "Septiembre": {...}} (en ese orden).
    combinadas: {"Octubre": ["D21:E21"]} rangos combinados por hoja.
    filas_ocultas: {"Octubre": [3]} filas ocultas por hoja.
    """
    libro = Workbook()
    libro.remove(libro.active)  # openpyxl crea una hoja "Sheet" que no queremos
    for nombre, celdas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for coordenada, valor in celdas.items():
            hoja[coordenada] = valor
        for rango in (combinadas or {}).get(nombre, []):
            hoja.merge_cells(rango)
        for fila in (filas_ocultas or {}).get(nombre, []):
            hoja.row_dimensions[fila].hidden = True
    return libro


def a_bytes(libro):
    """Guarda el Workbook en memoria y devuelve el contenido del .xlsx."""
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def xlsx(celdas, nombre_hoja="Octubre", **opciones):
    """Atajo: un .xlsx de una sola hoja a partir de un diccionario de celdas."""
    return a_bytes(armar_libro({nombre_hoja: celdas}, **opciones))


def planilla_prueba():
    """Bytes de `planilla-prueba.xlsx` tal cual (para variantes, ver planilla_con)."""
    return xlsx(celdas_planilla())


def planilla_con(cambios=None, borrar=()):
    """`planilla-prueba.xlsx` con celdas cambiadas/agregadas y otras borradas.

    cambios: {"E2": "Pago"}; borrar: ["C14", ...] (celdas que quedan vacías).
    """
    celdas = celdas_planilla()
    celdas.update(cambios or {})
    for coordenada in borrar:
        celdas.pop(coordenada, None)
    return xlsx(celdas)


def _separar(coordenada):
    letras = re.match(r"([A-Z]+)(\d+)", coordenada)
    return letras.group(1), int(letras.group(2))


def insertar_fila(celdas, despues_de, nuevas=None):
    """Inserta una fila del Excel nueva debajo de `despues_de`: las filas de abajo bajan una.

    nuevas: {"D": "Concepto", "E": "Importe"} (columna -> valor); None = fila vacía.
    """
    resultado = {}
    for coordenada, valor in celdas.items():
        columna, fila = _separar(coordenada)
        resultado[f"{columna}{fila + 1 if fila > despues_de else fila}"] = valor
    for columna, valor in (nuevas or {}).items():
        resultado[f"{columna}{despues_de + 1}"] = valor
    return resultado


def con_valores_guardados(contenido, valores):
    """Simula un archivo guardado por Excel: pone el valor calculado en las fórmulas.

    openpyxl guarda las fórmulas SIN valor (`<f>..</f><v />`). Excel guarda también el
    resultado. Esta función completa ese `<v />` a mano en el XML de la primera hoja.
    valores: {"D3": ("n", "300")}; tipos: "n" número, "str" texto, "e" error, "b" booleano.
    No es un archivo real de Excel: ver "pendiente" en el reporte (CA-52.5 / CA-58.4).
    """
    origen = zipfile.ZipFile(io.BytesIO(contenido))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for nombre in origen.namelist():
            datos = origen.read(nombre)
            if nombre == "xl/worksheets/sheet1.xml":
                texto = datos.decode("utf-8")
                for coordenada, (tipo, valor) in valores.items():
                    patron = re.compile(
                        rf'<c r="{coordenada}"([^>]*)>(<f>.*?</f>)<v\s*/></c>'
                    )
                    texto, cantidad = patron.subn(
                        rf'<c r="{coordenada}"\1 t="{tipo}">\2<v>{valor}</v></c>', texto
                    )
                    assert cantidad == 1, f"no se encontró la fórmula en {coordenada}"
                datos = texto.encode("utf-8")
            destino.writestr(nombre, datos)
    return salida.getvalue()


def reemplazar_en_zip(contenido, nombre_interno, nuevo):
    """Devuelve el .xlsx con un archivo interno reemplazado (para armar archivos dañados)."""
    origen = zipfile.ZipFile(io.BytesIO(contenido))
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
        for nombre in origen.namelist():
            destino.writestr(
                nombre, nuevo if nombre == nombre_interno else origen.read(nombre)
            )
    return salida.getvalue()


def resumen_filas(hoja):
    """Lista de tuplas (fila_excel, columna, nombre, monto_cent) para comparar fácil."""
    return [(f.fila_excel, f.columna, f.nombre, f.monto_cent) for f in hoja.filas]
